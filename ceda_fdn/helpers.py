"""UI-independent helpers around the two CEDA engines (the engines themselves are kept verbatim).
Nothing here changes the calculation method; it only validates input, ranks options and verifies output.
"""
import math, os, shutil, subprocess, tempfile, glob

from . import pile_capacity as pc

COMPETENT_SU = 10.0   # t/m2 – "competent layer" threshold used by the Step 01 skill (Engineering judgment)


# ------------------------------------------------------------------ Step 01
def layers_to_boreholes(rows, gwl_by_bh):
    """rows: list of dicts with BH, name, type, top, bot, gamma, N, Su.  Returns (boreholes, errors)."""
    errs, bhs = [], {}
    for i, r in enumerate(rows, 1):
        bh = str(r.get("BH") or "").strip()
        if not bh:
            continue
        try:
            L = {"name": str(r.get("name") or "").strip() or "-", "type": str(r.get("type") or "").strip().capitalize(),
                 "top": float(r["top"]), "bot": float(r["bot"]), "gamma": float(r["gamma"]),
                 "N": float(r.get("N") or 0), "Su": float(r.get("Su") or 0)}
        except (TypeError, ValueError, KeyError):
            errs.append(f"แถว {i} ({bh}): top/bot/γt ต้องเป็นตัวเลข"); continue
        if L["type"] not in ("Clay", "Sand"):
            errs.append(f"แถว {i} ({bh}): type ต้องเป็น Clay หรือ Sand")
        if L["bot"] <= L["top"]:
            errs.append(f"แถว {i} ({bh}): bot ต้องมากกว่า top")
        if L["gamma"] <= 0:
            errs.append(f"แถว {i} ({bh}): γt ต้อง > 0")
        bhs.setdefault(bh, []).append(L)
    out = []
    for name, layers in bhs.items():
        layers.sort(key=lambda L: L["top"])
        if abs(layers[0]["top"]) > 1e-6:
            errs.append(f"{name}: ชั้นแรกต้องเริ่มที่ 0.0 m")
        for a, b in zip(layers, layers[1:]):
            if abs(a["bot"] - b["top"]) > 1e-6:
                errs.append(f"{name}: ชั้นไม่ต่อเนื่องที่ {a['bot']} / {b['top']} m")
        g = gwl_by_bh.get(name)
        if g is None:
            errs.append(f"{name}: ยังไม่ได้ระบุ GWL")
        out.append({"name": name, "gwl": g, "layers": layers})
    if not out:
        errs.append("ยังไม่มีข้อมูลชั้นดิน")
    return out, errs


def recommend(cfg, res):
    """One option per D at the governing (deepest) tip over all BHs, re-checked in every BH.
    Recommended = least concrete volume, tip competent in all BHs, soil-governed, Qa >= target in all BHs."""
    if not cfg.get("target_qa"):
        return []
    tq = cfg["target_qa"]; out = []
    for D in cfg["diameters"]:
        tips = []
        for bh in res["boreholes"]:
            r = next((x for x in bh["back_calc"] if x["D"] == D), None)
            tips.append(r.get("tip") if r else None)
        if any(t is None for t in tips):
            out.append({"D": D, "tip": None, "status": "ไม่ถึง target ในบาง BH"}); continue
        tip = max(tips)
        per = [pc.capacity(bh["props"], D, tip, cfg["cutoff"], cfg["fs_soil"], cfg["fc"], cfg["qstr_coef"])
               for bh in res["boreholes"]]
        competent = all(p["tip_type"] == "Sand" or p["tip_Su"] >= COMPETENT_SU for p in per)
        soil_gov = all(p["governs"] == "soil" for p in per)
        qa_min = min(p["Qa"] for p in per)
        vol = math.pi * D * D / 4 * (tip - cfg["cutoff"])
        out.append({"D": D, "tip": tip, "L": tip - cfg["cutoff"], "vol": vol, "Qa_min": qa_min,
                    "tip_layers": ", ".join(sorted({p["tip_layer"] for p in per})),
                    "competent": competent, "soil_gov": soil_gov, "Qa_ok": qa_min >= tq - 1e-9})
    ok = [o for o in out if o.get("tip") is not None and o["competent"] and o["soil_gov"] and o["Qa_ok"]]
    best = min(ok, key=lambda o: o["vol"]) if ok else None
    for o in out:
        o["recommended"] = o is best
    return out


# ------------------------------------------------------------------ Step 02
def cap_figure(r, D):
    """Plan view of one cap type: outline, piles, pedestal, critical sections."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Circle, Rectangle
    fig, ax = plt.subplots(figsize=(4.2, 4.2))
    xs, ys = zip(*r["outline"])
    ax.fill(xs, ys, fc="#e8eef5", ec="#1f4e79", lw=1.5)
    cx, cy = r["pedestal"]
    ax.add_patch(Rectangle((-cx / 2, -cy / 2), cx, cy, fc="#9bb7d4", ec="#1f4e79"))
    d = r["checks"]["punch_col"]["d"]
    ax.add_patch(Rectangle((-(cx + d) / 2, -(cy + d) / 2), cx + d, cy + d, fill=False, ls="--", ec="#c00000", lw=1))
    for i, (x, y) in enumerate(r["coords"], 1):
        ax.add_patch(Circle((x, y), D / 2, fc="white", ec="black"))
        ax.text(x, y, str(i), ha="center", va="center", fontsize=8)
    ax.set_aspect("equal"); ax.set_title(f"{r['name']}  h = {r['h']} cm", fontsize=10)
    ax.tick_params(labelsize=7); ax.grid(alpha=0.3)
    fig.tight_layout()
    return fig


def find_soffice():
    for c in [shutil.which("soffice"), shutil.which("soffice.exe"),
              r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]:
        if c and os.path.exists(c):
            return c
    return None


def verify_cap_excel(xlsx_path, res, tol=0.001):
    """Recalculate the Excel with LibreOffice and compare Mu, Vu, vu and punching vu with the engine (tol 0.1%).
    Returns (list of mismatch strings, error-cell strings) or raises RuntimeError if LibreOffice is missing."""
    so = find_soffice()
    if not so:
        raise RuntimeError("ไม่พบ LibreOffice (soffice) – ข้ามการตรวจ Excel อัตโนมัติ")
    from openpyxl import load_workbook
    tmp = tempfile.mkdtemp()
    subprocess.run([so, "--headless", "--convert-to", "xlsx", "--outdir", tmp, xlsx_path],
                   check=True, capture_output=True, timeout=180)
    wb = load_workbook(glob.glob(os.path.join(tmp, "*.xlsx"))[0], data_only=True)
    bad, errcells = [], []
    for r in res:
        ws = wb[r["name"][:30]]
        lab = {}
        for row in ws.iter_rows(min_col=1, max_col=3):
            a, c = row[0].value, row[2].value
            if isinstance(c, str) and c.startswith("#"):
                errcells.append(f"{ws.title}!{row[2].coordinate} = {c}")
            if isinstance(a, str):
                lab.setdefault(a, []).append(c)
        X, Y, P = r["checks"]["dir_x"], r["checks"]["dir_y"], r["checks"]["punch_col"]
        pairs = []
        for ax, C in (("X", X), ("Y", Y)):
            key = next(k for k in lab if k.startswith(f"Mu {ax} "))
            pairs.append((f"Mu {ax}", lab[key][0], C["Mu_tm"] * 1000))
        vus = lab.get("Vu = Pu·max(Σfrac+,Σfrac-)", [None, None])
        vs = lab.get("vu = Vu/(b d)", [None, None])
        pairs += [("Vu X", vus[0], X["Vu_t"] * 1000), ("Vu Y", vus[1], Y["Vu_t"] * 1000),
                  ("vu X", vs[0], X["vu"]), ("vu Y", vs[1], Y["vu"])]
        if r["n"] > 1:
            pairs.append(("punch vu", lab.get("vu = Vu/(b0 d)", [None])[0], P["vu"]))
        for nm, xl, en in pairs:
            if xl is None or not isinstance(xl, (int, float)):
                bad.append(f"{r['name']} {nm}: Excel = {xl}"); continue
            if abs(xl - en) > tol * max(abs(en), 1e-6) and abs(xl - en) > 1e-6:
                bad.append(f"{r['name']} {nm}: Excel {xl:,.3f} ≠ engine {en:,.3f}")
    shutil.rmtree(tmp, ignore_errors=True)
    return bad, errcells
