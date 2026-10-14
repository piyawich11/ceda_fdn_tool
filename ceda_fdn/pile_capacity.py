#!/usr/bin/env python3
"""CEDA bored-pile capacity engine (reproduces CEDA '01_01/01_02 Bore Pile Capacity' sheets).
Units: t, m, t/m2, ksc.  Input JSON:
{
 "project": "...", "slurry": "polymer"|"bentonite", "cutoff": 1.0, "fs_soil": 2.5,
 "fc": 280, "qstr_coef": 0.29,
 "diameters": [0.8, 1.0, 1.2], "tip_depths": [20, 25, ...] (optional, else scan),
 "scan": {"from": 10, "to": 60, "step": 0.5},
 "target_qa": 215,                       # optional -> back-calculation
 "boreholes": [{"name":"BH-1","gwl":1.55,
     "layers":[{"name":"CH","type":"Clay","top":0,"bot":3,"gamma":1.36,"N":0,"Su":0}, ...]}]
}
Method (CEDA sheet, kept as-is; deviations are flagged, never applied silently):
 clay : alpha = ROUNDDOWN(1.0297-0.2165 ln Su, 3)  (capped at 1.0 -> flagged C1),  fs = alpha*Su,  qb = 9 Su
 sand : CN = 0.77 log10(20/(sv'/10)), N' = CN*N, phi = -0.0018N'^2+0.3817N'+26.521
        beta = 0.05 e^(0.055phi) [polymer] | 0.0002 e^(0.2035phi) [bentonite], fs = beta*sv'
        Nq = ROUNDDOWN(0.778 e^(0.053phi),1), qb = Nq*sv'
 sv' at mid-layer (total gamma, minus gamma_w*(z-GWL)); qb uses sv' at mid of the tip layer (sheet method)
 Qult = Qs(cut-off..tip) + Qb ; Qa_soil = Qult/FS ; Qstr = 0.29 fc' Ag ; Qa = min(Qa_soil, Qstr)
"""
import json, math, sys

def rounddown(x, n):
    f = 10 ** n
    return math.floor(x * f) / f if x >= 0 else -math.floor(-x * f) / f

def layer_props(bh, slurry, warn):
    gwl = bh["gwl"]; out = []; acc = 0.0
    for i, L in enumerate(bh["layers"]):
        t, b, g = L["top"], L["bot"], L["gamma"]
        zmid = (t + b) / 2
        sv = acc + (b - t) / 2 * g - max(0.0, zmid - gwl) * 1.0
        acc += (b - t) * g
        typ = L["type"].capitalize(); Su = L.get("Su", 0) or 0; N = L.get("N", 0) or 0
        p = dict(L, zmid=zmid, sv=sv, alpha=0, CN=0, N1=0, phi=0, beta=0, Nq=0, fs=0, qb=0)
        if typ == "Clay":
            if Su <= 0 and N > 0:
                warn.append(f"{bh['name']} layer {i+1} ({t}-{b} m): Clay with N={N} but Su=0 -> friction=0 "
                            f"(sheet note suggests Su=0.68N={0.68*N:.1f} t/m2 for stiff clay; ask designer) [C2]")
            if Su > 0:
                a = rounddown(1.0297 - 0.2165 * math.log(Su), 3)
                if a > 1.0:
                    warn.append(f"{bh['name']} layer {i+1}: alpha={a:.3f}>1.0 capped to 1.0 [C1]"); a = 1.0
                p.update(alpha=a, fs=a * Su, qb=9 * Su)
        elif typ == "Sand":
            CN = 0.77 * math.log10(20 / (sv / 10)) if sv > 0 else 0
            if CN > 2.0:
                warn.append(f"{bh['name']} layer {i+1}: CN={CN:.2f}>2.0 (sheet does not cap; flagged only)")
            N1 = N * CN
            phi = -0.0018 * N1 ** 2 + 0.3817 * N1 + 26.521
            beta = 0.05 * math.exp(0.055 * phi) if slurry == "polymer" else 0.0002 * math.exp(0.2035 * phi)
            Nq = rounddown(0.778 * math.exp(0.053 * phi), 1)
            p.update(CN=CN, N1=N1, phi=phi, beta=beta, Nq=Nq, fs=beta * sv, qb=Nq * sv)
        out.append(p)
    return out

def capacity(props, D, tip, cutoff, FS, fc, qcoef):
    P = math.pi * D; A = math.pi * D * D / 4
    qs = 0.0; qb = 0.0; tip_layer = None
    for p in props:
        lo, hi = max(p["top"], cutoff), min(p["bot"], tip)
        if hi > lo: qs += p["fs"] * (hi - lo)
        if p["top"] <= tip < p["bot"]: qb = p["qb"]; tip_layer = p
    if tip_layer is None and props and abs(tip - props[-1]["bot"]) < 1e-9:
        tip_layer = props[-1]; qb = tip_layer["qb"]
    Qs, Qb = qs * P, qb * A
    Qult = Qs + Qb; Qsoil = Qult / FS; Qstr = qcoef * fc * A * 10
    return dict(D=D, tip=tip, L=tip - cutoff, Qs=Qs, Qb=Qb, Qult=Qult, Qa_soil=Qsoil, Qstr=Qstr,
                Qa=min(Qsoil, Qstr), governs="soil" if Qsoil <= Qstr else "structure",
                tip_layer=(tip_layer or {}).get("name", "?"), tip_type=(tip_layer or {}).get("type", "?"),
                tip_Su=(tip_layer or {}).get("Su", 0) or 0)

def run(cfg):
    warn = []; res = {"boreholes": [], "warnings": warn}
    slurry = cfg.get("slurry", "polymer").lower(); FS = cfg.get("fs_soil", 2.5)
    cutoff = cfg.get("cutoff", 1.0); fc = cfg.get("fc", 280); qc = cfg.get("qstr_coef", 0.29)
    for bh in cfg["boreholes"]:
        props = layer_props(bh, slurry, warn)
        maxd = props[-1]["bot"]
        if "tip_depths" in cfg: tips = cfg["tip_depths"]
        else:
            s = cfg.get("scan", {}); a, b, st = s.get("from", 10), min(s.get("to", maxd), maxd), s.get("step", 0.5)
            tips = [round(a + i * st, 3) for i in range(int(round((b - a) / st)) + 1)]
        rows = [capacity(props, D, t, cutoff, FS, fc, qc) for D in cfg["diameters"] for t in tips]
        out = {"name": bh["name"], "props": props, "rows": rows}
        if cfg.get("target_qa"):
            tq = cfg["target_qa"]; pick = []
            for D in cfg["diameters"]:
                ok = [r for r in rows if r["D"] == D and r["Qa"] >= tq]
                if not ok: pick.append({"D": D, "tip": None, "note": "target not reached (soil or Qstr)"}); continue
                r = min(ok, key=lambda r: r["tip"])
                r = dict(r, vol=math.pi * D * D / 4 * r["L"],
                         note="tip in Sand/stiff layer" if (r["tip_type"] == "Sand" or r["tip_Su"] >= 10)
                         else "CHECK: tip in soft/medium clay")
                pick.append(r)
            out["back_calc"] = pick
        res["boreholes"].append(out)
    if len(cfg["boreholes"]) > 1 and cfg.get("target_qa"):
        env = {}
        for bh in res["boreholes"]:
            for r in bh["back_calc"]:
                if r.get("tip") is None: env[r["D"]] = None; continue
                if r["D"] in env and env[r["D"]] is None: continue
                env[r["D"]] = max(env.get(r["D"]) or 0, r["tip"])
        res["governing_tip"] = env   # deepest required tip over all boreholes
    return res

def to_excel(cfg, res, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "Summary"
    H = Font(bold=True); fill = PatternFill("solid", fgColor="DDEBF7")
    ws.append([cfg.get("project", ""), "Bored pile capacity", cfg.get("slurry", "polymer")]); ws["A1"].font = H
    ws.append(["FS", cfg.get("fs_soil", 2.5), "cut-off (m)", cfg.get("cutoff", 1.0), "fc' (ksc)", cfg.get("fc", 280),
               "Qstr coef", cfg.get("qstr_coef", 0.29), "target Qa (t)", cfg.get("target_qa")])
    ws.append([])
    if cfg.get("target_qa"):
        ws.append(["Back-calculation: shortest tip with Qa >= target"]); ws.cell(ws.max_row, 1).font = H
        ws.append(["BH", "D (m)", "Tip (m)", "L (m)", "Qs (t)", "Qb (t)", "Qult (t)", "Qa soil", "Qstr", "Qa (t)", "governs", "tip layer", "note"])
        for c in ws[ws.max_row]: c.font = H; c.fill = fill
        for bh in res["boreholes"]:
            for r in bh["back_calc"]:
                ws.append([bh["name"], r["D"], r.get("tip"), r.get("L"), r.get("Qs"), r.get("Qb"), r.get("Qult"),
                           r.get("Qa_soil"), r.get("Qstr"), r.get("Qa"), r.get("governs"), r.get("tip_layer"), r.get("note")])
        if "governing_tip" in res:
            ws.append(["Governing (deepest over BHs)"] + [f"D{d}: {t}" for d, t in res["governing_tip"].items()])
        ws.append([])
    ws.append(["Warnings"]); ws.cell(ws.max_row, 1).font = H
    for w in res["warnings"]: ws.append([w])
    for bh in res["boreholes"]:
        s = wb.create_sheet(bh["name"][:28])
        s.append(["Layer", "Type", "Top", "Bot", "gamma", "N", "Su", "zmid", "sv'", "alpha", "CN", "N'", "phi", "beta", "Nq", "fs (t/m2)", "qb (t/m2)"])
        for c in s[1]: c.font = H; c.fill = fill
        for p in bh["props"]:
            s.append([p["name"], p["type"], p["top"], p["bot"], p["gamma"], p.get("N", 0), p.get("Su", 0), p["zmid"], round(p["sv"], 3),
                      p["alpha"], round(p["CN"], 3), round(p["N1"], 2), round(p["phi"], 2), round(p["beta"], 4), p["Nq"], round(p["fs"], 3), round(p["qb"], 2)])
        s.append([])
        s.append(["D (m)", "Tip (m)", "L (m)", "Qs (t)", "Qb (t)", "Qult (t)", "Qa soil (t)", "Qstr (t)", "Qa (t)", "governs", "tip layer"])
        for c in s[s.max_row]: c.font = H; c.fill = fill
        for r in bh["rows"]:
            s.append([r["D"], r["tip"], r["L"], round(r["Qs"], 1), round(r["Qb"], 1), round(r["Qult"], 1), round(r["Qa_soil"], 1),
                      round(r["Qstr"], 1), round(r["Qa"], 1), r["governs"], r["tip_layer"]])
    for s in wb.worksheets:
        for col in s.columns: s.column_dimensions[col[0].column_letter].width = 11
    wb.save(path)

if __name__ == "__main__":
    cfg = json.load(open(sys.argv[1])); res = run(cfg)
    if len(sys.argv) > 2: to_excel(cfg, res, sys.argv[2])
    json.dump(res, open(sys.argv[1].replace(".json", "_result.json"), "w"), indent=1, default=str)
    for w in res["warnings"]: print("WARN", w)
    for bh in res["boreholes"]:
        for r in bh.get("back_calc", []): print(bh["name"], r)
